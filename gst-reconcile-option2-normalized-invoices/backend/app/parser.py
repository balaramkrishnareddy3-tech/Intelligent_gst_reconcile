import csv
import io
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


def _clean_key(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


ALIASES = {
    "invoice_number": ["invoice_number", "invoice_no", "invoice", "invoice_id", "inv_no", "document_number", "doc_no", "inum"],
    "invoice_date": ["invoice_date", "date", "invoice_dt", "document_date", "dt"],
    "vendor_gstin": ["vendor_gstin", "supplier_gstin", "gstin", "gstin_no", "supplier_gst", "seller_gstin", "ctin"],
    "vendor_name": ["vendor_name", "supplier_name", "vendor", "supplier", "seller_name", "trade_name"],
    "taxable_value": ["taxable_value", "taxable", "taxable_amount", "taxablevalue", "txval", "subtotal", "sub_total"],
    "cgst": ["cgst", "cgst_amount", "gst_cgst"],
    "sgst": ["sgst", "sgst_amount", "gst_sgst"],
    "igst": ["igst", "igst_amount", "gst_igst"],
    "total_gst": ["total_gst", "gst", "total_tax", "tax_amount", "tax", "totalgst"],
    "total_value": ["total_value", "invoice_value", "invoice_amount", "gross_value", "total", "total_due", "val"],
}


def _number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = text.replace(",", "").replace("₹", "").replace("$", "").replace("€", "").replace("£", "")
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text:
        return 0.0
    try:
        number = float(text)
        return -number if negative else number
    except ValueError:
        return 0.0


def _date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip().replace(",", "")
    for fmt in (
        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d/%b/%Y",
        "%d %b %Y", "%B %d %Y", "%b %d %Y", "%d %B %Y",
    ):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    match = re.search(r"\b(\d{4})[-/](\d{1,2})[-/](\d{1,2})\b", text)
    if match:
        try:
            return date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()
        except ValueError:
            pass
    return text[:10] if text else datetime.now().date().isoformat()


def _pick(row: dict[str, Any], field: str, default: Any = None) -> Any:
    normalized = {_clean_key(k): v for k, v in row.items() if k is not None}
    for alias in ALIASES[field]:
        if _clean_key(alias) in normalized:
            return normalized[_clean_key(alias)]
    return default


def _rows_from_csv(content: bytes) -> tuple[list[dict[str, Any]], list[list[str]]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel
    grid = [[str(cell or "").strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect=dialect)]
    while grid and not any(grid[-1]):
        grid.pop()
    if not grid:
        return [], []
    headers = grid[0]
    normalized_headers = {_clean_key(h) for h in headers if h}
    known_headers = {alias for aliases in ALIASES.values() for alias in map(_clean_key, aliases)}
    is_structured = len(normalized_headers & known_headers) >= 2
    if not is_structured:
        return [], grid
    return [dict(zip(headers, row)) for row in grid[1:] if any(row)], grid


def _rows_from_xlsx(content: bytes) -> tuple[list[dict[str, Any]], list[list[str]]]:
    if load_workbook is None:
        raise ValueError("Excel support is not installed. Run pip install -r requirements.txt.")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return [], []
    grid = [[str(x).strip() if x is not None else "" for x in row] for row in values]
    headers = grid[0]
    normalized_headers = {_clean_key(h) for h in headers if h}
    known_headers = {alias for aliases in ALIASES.values() for alias in map(_clean_key, aliases)}
    is_structured = len(normalized_headers & known_headers) >= 2
    if not is_structured:
        return [], grid
    return [dict(zip(headers, row)) for row in grid[1:] if any(row)], grid


def _rows_from_json(content: bytes) -> list[dict[str, Any]]:
    data = json.loads(content.decode("utf-8-sig"))
    if isinstance(data, dict):
        for key in ("invoices", "records", "data", "items"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    if not isinstance(data, list):
        raise ValueError("JSON must contain an object or an array of invoice records.")
    return [r for r in data if isinstance(r, dict)]


def _all_cells(grid: list[list[str]]) -> list[str]:
    return [cell.strip() for row in grid for cell in row if cell and cell.strip()]


def _joined(grid: list[list[str]]) -> str:
    return "\n".join(" | ".join(cell for cell in row if cell) for row in grid if any(row))


def _find_label_value(grid: list[list[str]], labels: tuple[str, ...]) -> str:
    for row in grid:
        for idx, cell in enumerate(row):
            clean = re.sub(r"\s+", " ", cell.strip().lower())
            if any(clean.startswith(label) for label in labels):
                # Prefer the value in the next column.
                for nxt in row[idx + 1:]:
                    if nxt.strip():
                        return nxt.strip()
                # Or text after ':' / '#'.
                if ":" in cell:
                    value = cell.split(":", 1)[1].strip()
                    if value:
                        return value
                if "#" in cell:
                    value = cell.split("#", 1)[1].strip()
                    if value:
                        return value
    return ""


def _find_amount_after_label(grid: list[list[str]], patterns: tuple[str, ...]) -> float:
    for row in grid:
        for idx, cell in enumerate(row):
            clean = re.sub(r"\s+", " ", cell.strip().lower())
            if any(re.search(pattern, clean) for pattern in patterns):
                # Search same row to the right first.
                for nxt in reversed(row[idx + 1:]):
                    if re.search(r"[-+]?[$₹€£]?\s*\(?[0-9][0-9,]*(?:\.[0-9]+)?\)?", nxt):
                        return _number(nxt)
                # Then parse an amount after the label in the same cell.
                after = re.split(r":", cell, maxsplit=1)[-1]
                match = re.search(r"[-+]?[$₹€£]?\s*\(?[0-9][0-9,]*(?:\.[0-9]+)?\)?", after)
                if match:
                    return _number(match.group(0))
    return 0.0


def _parse_unstructured_grid(grid: list[list[str]]) -> dict[str, Any] | None:
    text = _joined(grid)
    lower = text.lower()
    if "invoice" not in lower:
        return None

    labeled_invoice = _find_label_value(grid, ("invoice #", "invoice no", "invoice number"))
    invoice_match = re.search(
        r"(?:invoice\s*(?:#|no\.?|number)|inv(?:oice)?\s*(?:#|no\.?))\s*[:#-]?\s*([A-Z0-9][A-Z0-9/_-]{2,})",
        text,
        flags=re.IGNORECASE,
    )
    invoice_number = str(labeled_invoice or (invoice_match.group(1).strip() if invoice_match else ""))
    invoice_number = str(invoice_number or "").strip()
    if not invoice_number or invoice_number.lower() in {"invoice", "number"}:
        return None

    date_value = _find_label_value(grid, ("date issued", "invoice date", "date"))
    if not date_value:
        date_match = re.search(
            r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},?\s+\d{4}\b",
            text,
            flags=re.IGNORECASE,
        )
        date_value = date_match.group(0) if date_match else ""

    gstin_match = re.search(r"\b\d{2}[A-Z0-9]{13}\b", text.upper())
    vendor_gstin = gstin_match.group(0) if gstin_match else ""

    # Prefer explicit supplier/vendor/from labels. Otherwise use the first meaningful
    # business-like line before the INVOICE marker (common in invoice CSV exports).
    vendor_name = _find_label_value(grid, ("supplier", "vendor", "seller", "from"))
    if not vendor_name:
        for row in grid:
            values = [x.strip() for x in row if x.strip()]
            if not values:
                continue
            candidate = values[0]
            if candidate.upper() == "INVOICE":
                break
            if candidate.lower() in {"invoice", "date", "terms", "due date", "payment instructions"}:
                continue
            if "@" in candidate or re.search(r"\d{3}.*\d{3}.*\d{4}", candidate):
                continue
            if re.search(r"\b(?:street|avenue|road|suite|ca|tx|ny|zip|postal)\b", candidate, re.I):
                continue
            if len(candidate) > 2:
                vendor_name = candidate
                break

    taxable = _find_amount_after_label(grid, (r"^subtotal$", r"^sub\s*total$", r"taxable"))
    total_tax = _find_amount_after_label(grid, (r"tax\s*\(?", r"total\s+tax", r"gst"))
    cgst = _find_amount_after_label(grid, (r"^cgst",))
    sgst = _find_amount_after_label(grid, (r"^sgst",))
    igst = _find_amount_after_label(grid, (r"^igst",))
    total_value = _find_amount_after_label(grid, (r"^total\s*(?:due|amount|value)$", r"^grand\s*total$"))

    # If there is no subtotal label, sum the Amount column from a line-item table.
    if taxable == 0:
        for row_index, row in enumerate(grid):
            headers = [re.sub(r"\s+", " ", c.lower()) for c in row]
            if "description" in headers and "amount" in headers:
                amount_idx = headers.index("amount")
                for item_row in grid[row_index + 1:]:
                    if any(re.search(r"subtotal|total due|tax", c, re.I) for c in item_row):
                        break
                    if amount_idx < len(item_row):
                        taxable += _number(item_row[amount_idx])
                break

    components = cgst + sgst + igst
    if total_tax == 0:
        total_tax = components
    if total_value == 0:
        total_value = taxable + total_tax

    return {
        "invoice_number": invoice_number,
        "invoice_date": _date(date_value),
        "vendor_gstin": vendor_gstin,
        "vendor_name": vendor_name.strip(),
        "taxable_value": taxable,
        "cgst": cgst,
        "sgst": sgst,
        "igst": igst,
        "total_gst": total_tax,
        "total_value": total_value,
    }


def _normalize_structured(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for row in rows:
        invoice_number = str(_pick(row, "invoice_number", "") or "").strip()
        if not invoice_number:
            continue
        taxable = _number(_pick(row, "taxable_value"))
        cgst = _number(_pick(row, "cgst"))
        sgst = _number(_pick(row, "sgst"))
        igst = _number(_pick(row, "igst"))
        total_gst = _number(_pick(row, "total_gst"))
        if total_gst == 0:
            total_gst = cgst + sgst + igst
        total_value = _number(_pick(row, "total_value"))
        if total_value == 0:
            total_value = taxable + total_gst
        result.append({
            "invoice_number": invoice_number,
            "invoice_date": _date(_pick(row, "invoice_date")),
            "vendor_gstin": str(_pick(row, "vendor_gstin", "") or "").strip(),
            "vendor_name": str(_pick(row, "vendor_name", "") or "").strip(),
            "taxable_value": taxable,
            "cgst": cgst,
            "sgst": sgst,
            "igst": igst,
            "total_gst": total_gst,
            "total_value": total_value,
        })
    return result


def parse_invoice_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    ext = Path(filename).suffix.lower()
    grid: list[list[str]] = []
    if ext == ".csv":
        rows, grid = _rows_from_csv(content)
    elif ext in (".xlsx", ".xls"):
        if ext == ".xls":
            raise ValueError("Legacy .xls is not supported by the current parser. Save it as .xlsx or CSV and upload again.")
        rows, grid = _rows_from_xlsx(content)
    elif ext == ".json":
        rows = _rows_from_json(content)
    else:
        raise ValueError("Supported formats: CSV, XLSX, JSON.")

    result = _normalize_structured(rows) if rows else []
    if not result and grid:
        normalized = _parse_unstructured_grid(grid)
        if normalized:
            result = [normalized]

    if not result:
        raise ValueError(
            "No invoice records found. The parser accepts both standard GST columns and unstructured invoice exports "
            "with labels such as Invoice #, Date, Subtotal/Tax, and Total."
        )
    return result
